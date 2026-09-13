# -*- coding: utf-8 -*-
"""
凑单工具（通用版）

用法一（推荐给不熟悉命令行的朋友）：
    直接双击程序，弹出窗口后选择 Excel 表格即可。

用法二（命令行）：
    python coudan.py 输入表.xlsx [每单门槛=300] [输出表.xlsx] [搜索秒数=30]

输入表格式（取第一个工作表）：
    A 列 = 商品名称，B 列 = 单价，C 列 = 数量
    （单价可以是数字，也可以是公式；无表头或多行表头均可，非数字行自动跳过）

逻辑：
    1. 单数 = floor(总金额 / 门槛)，每单金额 >= 门槛 且尽量贴近门槛；
    2. 剩余不足门槛的商品单独列出；
    3. 价格相同的商品先合并求解，再按「需求大的单优先」拆回具体商品，使同品类尽量集中。
"""
import sys
import time
from collections import OrderedDict

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ============ 可调参数 ============
TIME_LIMIT = 30.0      # 搜索耗时上限（秒），越大结果越优
DEFAULT_THRESHOLD = 300.0
# =================================


def parse_num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(',', '')
    if s == '':
        return None
    try:
        return float(s)
    except ValueError:
        return None


def eval_formula(f, getcell):
    """兜底：解析 =D12/0.72 这类只含单元格引用与四则运算的简单公式。

    当 Excel 里没保存计算结果（data_only 取不到值）时才会用到；
    复杂函数（SUM、IF 等）解析不了，会返回 None。
    """
    if not isinstance(f, str) or not f.startswith('='):
        return None
    expr = f[1:]
    out, i = '', 0
    allowed = set('0123456789.+-*/() ')
    while i < len(expr):
        ch = expr[i]
        if ch.isalpha():
            j = i
            while j < len(expr) and expr[j].isalpha():
                j += 1
            k = j
            while k < len(expr) and expr[k].isdigit():
                k += 1
            if k == j:                       # 字母后面没跟数字，不是单元格引用
                return None
            col = 0
            for c in expr[i:j].upper():
                col = col * 26 + (ord(c) - 64)
            v = getcell(col, int(expr[j:k]))
            if v is None:
                return None
            out += repr(v)
            i = k
        elif ch in allowed:
            out += ch
            i += 1
        else:
            return None
    try:
        val = eval(out, {'__builtins__': {}}, {})
    except Exception:
        return None
    return float(val) if isinstance(val, (int, float)) else None


def read_items(path):
    """读取第一个工作表：A 列 = 商品名称，B 列 = 单价，C 列 = 数量"""
    wbf = load_workbook(path, data_only=False)
    wbv = load_workbook(path, data_only=True)
    wsf, wsv = wbf.worksheets[0], wbv.worksheets[0]

    def get_cell(col, row, depth=0):
        v = parse_num(wsv.cell(row, col).value)
        if v is not None:
            return v
        raw = wsf.cell(row, col).value
        if isinstance(raw, str) and raw.startswith('=') and depth < 5:
            return eval_formula(raw, lambda c, r: get_cell(c, r, depth + 1))
        return parse_num(raw)

    items = []
    for r in range(1, wsf.max_row + 1):
        name = wsf.cell(r, 1).value                    # A 列：商品名称
        if name is None or str(name).strip() == '':
            continue
        q = get_cell(3, r)                             # C 列：数量
        if q is None:
            continue
        q = int(round(q))
        if q <= 0:
            continue
        price = get_cell(2, r)                         # B 列：单价
        if price is None or price <= 0:
            continue
        items.append({'name': str(name).strip(), 'price': price, 'qty': q})
    return items


def solve(items, threshold, time_limit):
    """返回 (orders, leftover)；orders = [ {价格分: 件数} ]，leftover = {价格分: 件数}"""
    T = int(round(threshold * 100))
    total = sum(int(round(it['price'] * 100)) * it['qty'] for it in items)
    nbin = int(total // T) if T > 0 else 0
    if nbin <= 0:
        left = {}
        for it in items:
            left[int(round(it['price'] * 100))] = left.get(int(round(it['price'] * 100)), 0) + it['qty']
        return [], left

    gmap = OrderedDict()
    for it in items:
        gmap.setdefault(int(round(it['price'] * 100)), 0)
        gmap[int(round(it['price'] * 100))] += it['qty']
    prices = sorted(gmap.keys(), reverse=True)
    K = len(prices)
    counts0 = [gmap[p] for p in prices]
    budget = total - nbin * T          # 全部用完时的总超出上限
    best = [budget + 1, None]
    start = [time.time()]
    stop = [False]

    def combos(counts, s, lo, hi, cur, out, sel):
        if cur >= lo:
            out.append((tuple(sel), cur))
        if s >= K or cur > hi:
            return
        for i in range(s, K):
            if counts[i] <= 0 or cur + prices[i] > hi:
                continue
            c = 1
            while c <= counts[i]:
                if cur + prices[i] * c > hi:
                    break
                sel[i] = c
                combos(counts, i + 1, lo, hi, cur + prices[i] * c, out, sel)
                sel[i] = 0
                c += 1

    def dfs(bins_left, counts, over, acc):
        if stop[0]:
            return
        if bins_left % 3 == 0 and time.time() - start[0] > time_limit:
            stop[0] = True
            return
        if bins_left == 0:
            if over < best[0]:
                best[0] = over
                best[1] = [dict(a) for a in acc]
                if over == 0:
                    stop[0] = True
            return
        rem = sum(p * c for p, c in zip(prices, counts))
        if rem < T * bins_left or over >= best[0] or over > budget:
            return
        i0 = min(i for i in range(K) if counts[i] > 0)
        hi = T + min(budget, best[0] - 1) - over
        for c0 in range(1, counts[i0] + 1):
            base = prices[i0] * c0
            if base >= T:
                if base > hi:
                    break
                nc = list(counts); nc[i0] -= c0
                acc.append({i0: c0})
                dfs(bins_left - 1, nc, over + base - T, acc)
                acc.pop()
                continue
            sub, sel = [], [0] * K
            combos(counts, i0 + 1, T - base, hi - base, 0, sub, sel)
            dedup = {}
            for cc, v in sub:
                if cc not in dedup or v < dedup[cc]:
                    dedup[cc] = v
            for cc, v in sorted(dedup.items(), key=lambda t: t[1]):
                nc = list(counts); nc[i0] -= c0
                for i in range(K):
                    nc[i] -= cc[i]
                d = {i0: c0}
                for i in range(K):
                    if cc[i]:
                        d[i] = cc[i]
                acc.append(d)
                dfs(bins_left - 1, nc, over + base + v - T, acc)
                acc.pop()

    dfs(nbin, counts0, 0, [])

    if best[1] is None:                        # 兜底：贪心构造
        orders, left = greedy(prices, counts0, T, nbin)
        return orders, left
    orders = [{prices[i]: c for i, c in a.items() if c} for a in best[1]]
    used = {}
    for o in orders:
        for p, c in o.items():
            used[p] = used.get(p, 0) + c
    left = {p: counts0[i] - used.get(p, 0) for i, p in enumerate(prices)}
    left = {p: c for p, c in left.items() if c}
    return orders, left


def greedy(prices, counts, T, nbin):
    counts = list(counts)
    orders = []
    for _ in range(nbin):
        cur, binv = 0, {}
        for i in range(len(prices)):
            while counts[i] > 0 and cur < T:
                need = T - cur
                take = min(counts[i], max(1, -(-need // prices[i])))
                counts[i] -= take
                binv[prices[i]] = binv.get(prices[i], 0) + take
                cur += prices[i] * take
        if cur < T:                      # 凑不够门槛，回滚本单
            for p, c in binv.items():
                counts[prices.index(p)] += c
            break
        orders.append(binv)
    left = {p: counts[i] for i, p in enumerate(prices) if counts[i] > 0}
    return [o for o in orders if o], left


def build_output(items, orders, leftover, threshold, out_path):
    info = {it['name']: it['price'] for it in items}

    # 价格 -> 该价位的商品列表（保持原表顺序）
    by_price = OrderedDict()
    for it in items:
        by_price.setdefault(int(round(it['price'] * 100)), []).append(it)

    # 拆回具体商品
    final_orders = [{} for _ in orders]
    for p, group in by_price.items():
        need = {j: 0 for j in range(len(orders))}
        for j, o in enumerate(orders):
            need[j] = o.get(p, 0)
        for it in group:
            q = it['qty']
            for j in sorted(need, key=lambda x: -need[x]):
                if q <= 0:
                    break
                if need[j] <= 0:
                    continue
                take = min(q, need[j])
                need[j] -= take
                q -= take
                final_orders[j][it['name']] = final_orders[j].get(it['name'], 0) + take

    # 排序：让内容相同的单相邻
    def key(fo):
        return (sum(info[n] * c for n, c in fo.items()), tuple(sorted(fo.items())))
    order_idx = sorted(range(len(final_orders)), key=lambda j: key(final_orders[j]))
    final_orders = [final_orders[j] for j in order_idx]

    wb = Workbook()
    ws = wb.active
    ws.title = '凑单方案'
    head_fill = PatternFill('solid', start_color='D9E2F3')
    tot_fill = PatternFill('solid', start_color='FFF2CC')
    left_fill = PatternFill('solid', start_color='FCE4D6')
    bold = Font(name='Arial', bold=True)
    fnt = Font(name='Arial')
    thin = Side(style='thin', color='BFBFBF')
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws['A1'] = f'凑单方案（每单 ≥ {threshold:g} 元，尽量贴近 {threshold:g}）'
    ws['A1'].font = Font(name='Arial', bold=True, size=12)
    ws.merge_cells('A1:E1')
    for c, h in enumerate(['订单号', '商品名称', '单价(元)', '数量', '小计(元)'], 1):
        cell = ws.cell(2, c, h)
        cell.font = bold
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = bd

    row = 3
    for idx, fo in enumerate(final_orders, 1):
        start = row
        for nm, q in sorted(fo.items(), key=lambda t: -info[t[0]]):
            for c, v in enumerate([idx, nm, info[nm], q, info[nm] * q], 1):
                ws.cell(row, c, v).font = fnt
                ws.cell(row, c).border = bd
            row += 1
        ws.cell(row, 1, idx).font = bold
        ws.cell(row, 2, '订单合计').font = bold
        ws.cell(row, 4, sum(fo.values())).font = bold
        ws.cell(row, 5, sum(info[n] * c for n, c in fo.items())).font = bold
        for c in range(1, 6):
            ws.cell(row, c).border = bd
            ws.cell(row, c).fill = tot_fill
        row += 2

    left_total = 0
    # 剩余按具体商品分摊
    left_items = []
    for p, group in by_price.items():
        need = leftover.get(p, 0)
        for it in group:
            take = min(it['qty'], need)
            if take > 0:
                left_items.append((it['name'], info[it['name']], take))
                left_total += info[it['name']] * take
                need -= take
            if need <= 0:
                break
    if left_items:
        ws.cell(row, 1, f'以下不足 {threshold:g} 元，需单独下单').font = Font(name='Arial', bold=True, color='C00000')
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 1
        for nm, pr, q in left_items:
            for c, v in enumerate(['剩余', nm, pr, q, pr * q], 1):
                ws.cell(row, c, v).font = fnt
                ws.cell(row, c).border = bd
            row += 1
        ws.cell(row, 2, f'小计（不足 {threshold:g}）').font = bold
        ws.cell(row, 5, left_total).font = bold
        for c in range(1, 6):
            ws.cell(row, c).border = bd
            ws.cell(row, c).fill = left_fill

    for col, w in zip('ABCDE', [9, 26, 10, 8, 10]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A3'
    wb.save(out_path)
    return final_orders, left_items, left_total


def run(src, thr=DEFAULT_THRESHOLD, out=None, lim=TIME_LIMIT, verbose=True):
    """执行一次凑单，返回 (结果说明文本, 输出文件路径)"""
    if out is None:
        out = src.rsplit('.', 1)[0] + '_凑单结果.xlsx'
    items = read_items(src)
    if not items:
        raise ValueError('未读取到数据。请确认：第一个工作表的 A 列是商品名称、'
                         'B 列是单价、C 列是数量（数字）。')
    pmap = {it['name']: it['price'] for it in items}
    total = sum(it['price'] * it['qty'] for it in items)
    t0 = time.time()
    orders, leftover = solve(items, thr, lim)
    fo, left_items, left_total = build_output(items, orders, leftover, thr, out)

    lines = [f'共 {len(items)} 个商品，总金额 {total:.2f} 元',
             f'门槛 {thr:g} 元：{len(fo)} 单 + 剩余 {left_total:.2f} 元', '']
    for i, o in enumerate(fo, 1):
        s = sum(pmap[n] * c for n, c in o.items())
        lines.append(f'第 {i} 单（{s:.2f} 元）：' + '、'.join(f'{n} x{c}' for n, c in o.items()))
    if left_items:
        lines.append('')
        lines.append('剩余（不足门槛，需单独下单）：' +
                     '、'.join(f'{n} x{c}' for n, _, c in left_items))
    lines.append('')
    lines.append(f'结果已保存到：{out}')
    lines.append(f'（用时 {time.time()-t0:.1f} 秒）')
    text = '\n'.join(lines)
    if verbose:
        print(text)
    return text, out


def main():
    args = sys.argv[1:]
    if args:                                   # 命令行模式
        src = args[0]
        thr = float(args[1]) if len(args) > 1 else DEFAULT_THRESHOLD
        out = args[2] if len(args) > 2 else None
        lim = float(args[3]) if len(args) > 3 else TIME_LIMIT
        try:
            run(src, thr, out, lim)
        except Exception as e:
            print('出错：' + str(e))
        return
    # 没有参数：图形界面模式（双击打包版时走这里）
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, simpledialog
    except Exception:
        print(__doc__)
        return
    root = tk.Tk()
    root.withdraw()
    src = filedialog.askopenfilename(
        title='请选择商品表（Excel 文件）',
        filetypes=[('Excel 表格', '*.xlsx *.xlsm'), ('所有文件', '*.*')])
    if not src:
        return
    ans = simpledialog.askstring('每单最低金额', '请输入每单最低金额（默认 300）：',
                                 initialvalue='300', parent=root)
    if ans is None:
        return
    try:
        thr = float(ans)
    except ValueError:
        thr = DEFAULT_THRESHOLD
    messagebox.showinfo('开始凑单', '正在计算，请稍候（最多约 30 秒），完成后会弹出结果。')
    try:
        text, out = run(src, thr, None, TIME_LIMIT, verbose=False)
        messagebox.showinfo('凑单完成', text)
    except Exception as e:
        messagebox.showerror('出错了', str(e))
    finally:
        root.destroy()


if __name__ == '__main__':
    main()
